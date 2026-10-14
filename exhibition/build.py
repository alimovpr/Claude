"""Собирает contacts.xlsx из contacts.json (контакты с выставки, сгруппированы по странам)."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
COLUMNS = [
    ("country", "Страна производства", 14),
    ("category", "Тип компании", 18),
    ("city", "Город / регион", 22),
    ("company", "Компания", 24),
    ("booth", "Выставка / стенд", 20),
    ("website", "Сайт", 18),
    ("catalog", "Каталог", 28),
    ("profile", "Что производят", 45),
    ("contact_name", "Контакт (по визитке)", 18),
    ("introduced_as", "Как представился", 18),
    ("position", "Должность", 18),
    ("email", "Email", 24),
    ("phone", "Телефон", 16),
    ("address", "Адрес / офисы", 36),
    ("card_notes", "Пометки на визитке", 22),
    ("discussed", "Что обсуждали / интерес", 30),
    ("web_info", "Найдено в интернете", 50),
    ("sources", "Источники", 30),
    ("next_step", "Следующий шаг", 32),
    ("priority", "Приоритет", 11),
    ("date", "Дата", 11),
    ("photo", "Фото", 18),
]
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
COUNTRY_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main():
    rows = json.loads((HERE / "contacts.json").read_text(encoding="utf-8"))
    rows.sort(key=lambda r: (r["country"], r["company"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Контакты"
    ws.append(["№"] + [c[1] for c in COLUMNS])
    for i, r in enumerate(rows, 1):
        ws.append([i] + [r.get(c[0], "") for c in COLUMNS])

    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        for key in ("website", "catalog"):
            cell = row[1 + [c[0] for c in COLUMNS].index(key)]
            url = str(cell.value or "").split(" ")[0]
            if url.startswith("http") or url.startswith("www."):
                cell.hyperlink = url if url.startswith("http") else "https://" + url
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        row[1].fill = COUNTRY_FILL
        row[1].font = Font(name=FONT, size=10, bold=True)
    ws.column_dimensions["A"].width = 5
    for idx, (_, _, width) in enumerate(COLUMNS, 2):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # Сводка по странам — считается формулами от листа «Контакты»
    sm = wb.create_sheet("По странам")
    sm.append(["Страна производства", "Компаний", "Высокий приоритет"])
    last = len(rows) + 1
    pr = get_column_letter(2 + [c[0] for c in COLUMNS].index("priority"))
    for country in sorted({r["country"] for r in rows}):
        n = sm.max_row + 1
        sm.append([
            country,
            f"=COUNTIF(Контакты!$B$2:$B${last},A{n})",
            f"=COUNTIFS(Контакты!$B$2:$B${last},A{n},Контакты!${pr}$2:${pr}${last},\"Высокий\")",
        ])
    total = sm.max_row + 1
    sm.append(["Итого", f"=SUM(B2:B{total - 1})", f"=SUM(C2:C{total - 1})"])

    # Сводка по типу компании
    sm.append([])
    cat_header = sm.max_row + 1
    sm.append(["Тип компании", "Компаний"])
    for category in sorted({r["category"] for r in rows}):
        n = sm.max_row + 1
        sm.append([category, f"=COUNTIF(Контакты!$C$2:$C${last},A{n})"])
    for row in sm.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=row[0].row == total)
    for header_row in (1, cat_header):
        for cell in sm[header_row]:
            if cell.value:
                cell.font = Font(name=FONT, bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
    for col, width in zip("ABC", (28, 12, 20)):
        sm.column_dimensions[col].width = width

    wb.save(HERE / "contacts.xlsx")


if __name__ == "__main__":
    main()
